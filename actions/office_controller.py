# office_controller.py — KITT Office Expert
# Handles Word, Excel, and PowerPoint with full professional capability
# Dependencies: python-docx, openpyxl, python-pptx (pip install python-docx openpyxl python-pptx)

import os
import re
import sys
import json
import subprocess
from pathlib import Path
from datetime import datetime

# ── Word ────────────────────────────────────────────────────────────────────
try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    _DOCX = True
except ImportError:
    _DOCX = False

# ── Excel ────────────────────────────────────────────────────────────────────
try:
    import openpyxl
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.utils import get_column_letter
    _XLSX = True
except ImportError:
    _XLSX = False

# ── PowerPoint ───────────────────────────────────────────────────────────────
try:
    from pptx import Presentation
    from pptx.util import Inches as PInches, Pt as PPt, Emu
    from pptx.dml.color import RGBColor as PRGBColor
    from pptx.enum.text import PP_ALIGN
    _PPTX = True
except ImportError:
    _PPTX = False


def _desktop() -> Path:
    return Path.home() / "Desktop"

def _documents() -> Path:
    return Path.home() / "Documents"

def _resolve_path(path_str: str | None, default_name: str) -> Path:
    """Resolve a path string, defaulting to Desktop."""
    if not path_str:
        return _desktop() / default_name
    p = Path(path_str)
    if p.is_dir():
        return p / default_name
    return p

def _open_file(path: Path):
    """Open file with default application."""
    try:
        if sys.platform == "win32":
            os.startfile(str(path))
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])
    except Exception:
        pass


# ════════════════════════════════════════════════════════════════════════════
#  WORD
# ════════════════════════════════════════════════════════════════════════════

def _word_create(params: dict) -> str:
    if not _DOCX:
        return "Error: python-docx not installed. Run: pip install python-docx"

    title    = params.get("title", "Document")
    content  = params.get("content", "")
    style    = params.get("style", "professional")   # professional | academic | simple
    output   = _resolve_path(params.get("output_path"), f"{title}.docx")
    open_it  = params.get("open", True)

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3)
        section.right_margin  = Cm(2.5)

    # Title
    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.runs[0]
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    run.font.size = Pt(20)

    # Date line for professional style
    if style in ("professional", "academic"):
        date_p = doc.add_paragraph(datetime.now().strftime("%B %d, %Y"))
        date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_p.runs[0].font.color.rgb = RGBColor(0x70, 0x70, 0x70)
        date_p.runs[0].font.size = Pt(10)

    doc.add_paragraph()

    # Parse content sections
    if content:
        for block in content.split("\n\n"):
            block = block.strip()
            if not block:
                continue
            if block.startswith("## "):
                doc.add_heading(block[3:], level=2)
            elif block.startswith("# "):
                doc.add_heading(block[2:], level=1)
            elif block.startswith("- ") or block.startswith("* "):
                for line in block.splitlines():
                    line = line.lstrip("-* ").strip()
                    if line:
                        p = doc.add_paragraph(style="List Bullet")
                        p.add_run(line)
            elif re.match(r"^\d+\.", block.splitlines()[0]):
                for line in block.splitlines():
                    m = re.match(r"^\d+\.\s*(.*)", line)
                    if m:
                        p = doc.add_paragraph(style="List Number")
                        p.add_run(m.group(1))
            else:
                p = doc.add_paragraph(block)
                p.paragraph_format.space_after = Pt(6)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    if open_it:
        _open_file(output)
    return f"Word document created: {output.name}"


def _word_edit(params: dict) -> str:
    if not _DOCX:
        return "Error: python-docx not installed."
    file_path = params.get("file_path")
    if not file_path or not Path(file_path).exists():
        return "Error: File not found. Provide file_path."
    
    doc       = Document(file_path)
    action    = params.get("edit_action", "append")  # append | replace | add_heading | add_table
    content   = params.get("content", "")
    find_text = params.get("find", "")
    open_it   = params.get("open", True)

    if action == "append":
        doc.add_paragraph(content)
    elif action == "add_heading":
        level = int(params.get("level", 1))
        doc.add_heading(content, level=level)
    elif action == "replace" and find_text:
        for para in doc.paragraphs:
            if find_text in para.text:
                for run in para.runs:
                    run.text = run.text.replace(find_text, content)
    elif action == "add_table":
        rows = int(params.get("rows", 3))
        cols = int(params.get("cols", 3))
        table = doc.add_table(rows=rows, cols=cols)
        table.style = "Table Grid"
        headers = params.get("headers", [])
        data    = params.get("data", [])
        if headers:
            for i, h in enumerate(headers[:cols]):
                cell = table.rows[0].cells[i]
                cell.text = str(h)
                cell.paragraphs[0].runs[0].font.bold = True
        for r_idx, row_data in enumerate(data[:rows-1], start=1):
            for c_idx, val in enumerate(row_data[:cols]):
                table.rows[r_idx].cells[c_idx].text = str(val)

    doc.save(file_path)
    if open_it:
        _open_file(Path(file_path))
    return f"Word document updated: {Path(file_path).name}"


def _word_to_pdf(params: dict) -> str:
    file_path = params.get("file_path")
    if not file_path or not Path(file_path).exists():
        return "Error: File not found."
    try:
        if sys.platform == "win32":
            import comtypes.client
            word = comtypes.client.CreateObject("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(str(Path(file_path).resolve()))
            pdf_path = str(Path(file_path).with_suffix(".pdf"))
            doc.SaveAs(pdf_path, FileFormat=17)
            doc.Close()
            word.Quit()
            _open_file(Path(pdf_path))
            return f"Converted to PDF: {Path(pdf_path).name}"
        else:
            out_dir = str(Path(file_path).parent)
            subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf",
                           "--outdir", out_dir, file_path], check=True, capture_output=True)
            return f"Converted to PDF in {out_dir}"
    except Exception as e:
        return f"PDF conversion failed: {e}"


# ════════════════════════════════════════════════════════════════════════════
#  EXCEL
# ════════════════════════════════════════════════════════════════════════════

_HEADER_FILL = "1F497D"
_ALT_FILL    = "DCE6F1"

def _style_excel_header(ws, row: int, num_cols: int):
    """Apply professional header styling."""
    if not _XLSX:
        return
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font      = Font(bold=True, color="FFFFFF", size=11)
        cell.fill      = PatternFill("solid", fgColor=_HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = Border(
            bottom=Side(style="medium", color="FFFFFF"),
            right=Side(style="thin", color="FFFFFF")
        )

def _style_excel_data(ws, start_row: int, end_row: int, num_cols: int):
    """Apply alternating row colors."""
    if not _XLSX:
        return
    for row in range(start_row, end_row + 1):
        fill_color = _ALT_FILL if row % 2 == 0 else "FFFFFF"
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill      = PatternFill("solid", fgColor=fill_color)
            cell.alignment = Alignment(vertical="center")
            cell.border    = Border(
                bottom=Side(style="thin", color="CCCCCC"),
                right=Side(style="thin", color="CCCCCC")
            )

def _excel_create(params: dict) -> str:
    if not _XLSX:
        return "Error: openpyxl not installed. Run: pip install openpyxl"

    title    = params.get("title", "Spreadsheet")
    headers  = params.get("headers", [])
    data     = params.get("data", [])
    formulas = params.get("formulas", [])   # list of {"col": "B", "row": 5, "formula": "=SUM(B2:B4)"}
    add_chart= params.get("add_chart", False)
    chart_type = params.get("chart_type", "bar")   # bar | line | pie
    output   = _resolve_path(params.get("output_path"), f"{title}.xlsx")
    open_it  = params.get("open", True)

    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]

    # Title row
    ws.merge_cells(f"A1:{get_column_letter(max(len(headers), 1))}1")
    title_cell = ws["A1"]
    title_cell.value     = title
    title_cell.font      = Font(bold=True, size=14, color="FFFFFF")
    title_cell.fill      = PatternFill("solid", fgColor="2E4057")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Headers
    if headers:
        for col_idx, h in enumerate(headers, start=1):
            ws.cell(row=2, column=col_idx, value=str(h))
        _style_excel_header(ws, 2, len(headers))
        ws.row_dimensions[2].height = 22

    # Data rows
    data_start = 3
    for row_idx, row in enumerate(data, start=data_start):
        for col_idx, val in enumerate(row, start=1):
            # Auto-detect numbers
            try:
                ws.cell(row=row_idx, column=col_idx, value=float(val) if '.' in str(val) else int(val))
            except (ValueError, TypeError):
                ws.cell(row=row_idx, column=col_idx, value=val)
        ws.row_dimensions[row_idx].height = 18

    if data and headers:
        _style_excel_data(ws, data_start, data_start + len(data) - 1, len(headers))

    # Formulas
    for f in formulas:
        col  = f.get("col", "A")
        row  = f.get("row", data_start + len(data))
        expr = f.get("formula", "")
        if expr:
            cell = ws[f"{col}{row}"]
            cell.value = expr
            cell.font  = Font(bold=True, color="1F497D")

    # Auto-width columns
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

    # Freeze header row
    ws.freeze_panes = "A3"

    # Chart
    if add_chart and data and headers and len(headers) >= 2:
        data_ref = Reference(ws, min_col=2, min_row=2,
                             max_col=len(headers), max_row=2 + len(data))
        cats_ref = Reference(ws, min_col=1, min_row=3, max_row=2 + len(data))

        if chart_type == "pie":
            chart = PieChart()
            chart.title  = title
            chart.series.append(openpyxl.chart.Series(data_ref, title_from_data=True))
        elif chart_type == "line":
            chart = LineChart()
            chart.title  = title
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
        else:
            chart = BarChart()
            chart.title  = title
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)

        chart.width  = 18
        chart.height = 12
        anchor_row    = data_start + len(data) + 2
        ws.add_chart(chart, f"A{anchor_row}")

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    if open_it:
        _open_file(output)
    return f"Excel spreadsheet created: {output.name}"


def _excel_analyze(params: dict) -> str:
    if not _XLSX:
        return "Error: openpyxl not installed."
    file_path = params.get("file_path")
    if not file_path or not Path(file_path).exists():
        return "Error: File not found."

    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    summary = []
    summary.append(f"Sheets: {', '.join(wb.sheetnames)}")
    summary.append(f"Active sheet: {ws.title}")
    summary.append(f"Dimensions: {ws.dimensions}")

    # Read first rows
    rows = list(ws.iter_rows(max_row=6, values_only=True))
    if rows:
        summary.append(f"Headers: {list(rows[0])}")
        summary.append(f"Sample rows: {len(rows)-1}")
        
        # Basic numeric stats on first numeric column found
        for col_idx in range(len(rows[0])):
            vals = []
            for row in rows[1:]:
                try:
                    v = float(row[col_idx])
                    vals.append(v)
                except (TypeError, ValueError):
                    pass
            if vals:
                col_name = rows[0][col_idx] or f"Col{col_idx+1}"
                summary.append(f"Column '{col_name}': min={min(vals):.2f}, max={max(vals):.2f}, avg={sum(vals)/len(vals):.2f}")
                break

    wb.close()
    return " | ".join(summary)


def _excel_add_data(params: dict) -> str:
    if not _XLSX:
        return "Error: openpyxl not installed."
    file_path = params.get("file_path")
    if not file_path or not Path(file_path).exists():
        return "Error: File not found."

    wb   = load_workbook(file_path)
    ws   = wb.active
    data = params.get("data", [])
    
    # Find next empty row
    next_row = ws.max_row + 1
    for row_data in data:
        for col_idx, val in enumerate(row_data, start=1):
            try:
                ws.cell(row=next_row, column=col_idx,
                        value=float(val) if '.' in str(val) else int(val))
            except (ValueError, TypeError):
                ws.cell(row=next_row, column=col_idx, value=val)
        next_row += 1

    wb.save(file_path)
    _open_file(Path(file_path))
    return f"Added {len(data)} row(s) to {Path(file_path).name}"


# ════════════════════════════════════════════════════════════════════════════
#  POWERPOINT
# ════════════════════════════════════════════════════════════════════════════

def _pptx_create(params: dict) -> str:
    if not _PPTX:
        return "Error: python-pptx not installed. Run: pip install python-pptx"

    title   = params.get("title", "Presentation")
    slides  = params.get("slides", [])   # list of {"title": str, "content": str or list}
    theme   = params.get("theme", "dark")  # dark | light | blue
    output  = _resolve_path(params.get("output_path"), f"{title}.pptx")
    open_it = params.get("open", True)

    prs = Presentation()
    prs.slide_width  = PInches(13.33)
    prs.slide_height = PInches(7.5)

    # Theme colors
    themes = {
        "dark":  {"bg": "1A1A2E", "title": "E94560", "text": "EAEAEA", "accent": "16213E"},
        "blue":  {"bg": "1F497D", "title": "FFFFFF", "text": "DCE6F1", "accent": "2E75B6"},
        "light": {"bg": "FFFFFF", "title": "1F497D", "text": "333333", "accent": "DCE6F1"},
    }
    colors = themes.get(theme, themes["dark"])

    def hex_color(h: str) -> PRGBColor:
        return PRGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    def set_bg(slide, hex_str: str):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = hex_color(hex_str)

    # Title slide
    blank_layout = prs.slide_layouts[6]
    title_slide  = prs.slides.add_slide(blank_layout)
    set_bg(title_slide, colors["bg"])

    # Main title
    txb = title_slide.shapes.add_textbox(PInches(1), PInches(2.5), PInches(11.33), PInches(1.5))
    tf  = txb.text_frame
    tf.word_wrap = True
    p   = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.size  = PPt(44)
    run.font.bold  = True
    run.font.color.rgb = hex_color(colors["title"])

    # Subtitle (date)
    txb2 = title_slide.shapes.add_textbox(PInches(1), PInches(4.2), PInches(11.33), PInches(0.8))
    tf2  = txb2.text_frame
    p2   = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = datetime.now().strftime("%B %Y")
    run2.font.size  = PPt(20)
    run2.font.color.rgb = hex_color(colors["text"])

    # Content slides
    for slide_data in slides:
        s_title   = slide_data.get("title", "")
        s_content = slide_data.get("content", "")
        if isinstance(s_content, list):
            s_content = "\n".join(f"• {item}" for item in s_content)

        slide = prs.slides.add_slide(blank_layout)
        set_bg(slide, colors["bg"])

        # Accent bar top
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, PInches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = hex_color(colors["title"])
        bar.line.fill.background()

        # Slide title
        t_box = slide.shapes.add_textbox(PInches(0.5), PInches(0.3), PInches(12.33), PInches(1))
        t_tf  = t_box.text_frame
        t_p   = t_tf.paragraphs[0]
        t_run = t_p.add_run()
        t_run.text = s_title
        t_run.font.size  = PPt(28)
        t_run.font.bold  = True
        t_run.font.color.rgb = hex_color(colors["title"])

        # Content
        if s_content:
            c_box = slide.shapes.add_textbox(PInches(0.5), PInches(1.5), PInches(12.33), PInches(5.5))
            c_tf  = c_box.text_frame
            c_tf.word_wrap = True
            for i, line in enumerate(s_content.split("\n")):
                if i == 0:
                    p = c_tf.paragraphs[0]
                else:
                    p = c_tf.add_paragraph()
                p.space_before = PPt(4)
                run = p.add_run()
                run.text = line
                run.font.size  = PPt(18)
                run.font.color.rgb = hex_color(colors["text"])

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    if open_it:
        _open_file(output)
    return f"PowerPoint created with {len(slides)+1} slides: {output.name}"


# ════════════════════════════════════════════════════════════════════════════
#  MAIN DISPATCHER
# ════════════════════════════════════════════════════════════════════════════

def office_controller(parameters: dict, player=None, speak=None) -> str:
    action      = parameters.get("action", "").lower()
    app         = parameters.get("app", "word").lower()   # word | excel | powerpoint

    try:
        # ── WORD ───────────────────────────────────────────────────────────
        if app in ("word", "doc", "document"):
            if action in ("create", "write", "new"):
                return _word_create(parameters)
            elif action in ("edit", "update", "modify", "add"):
                return _word_edit(parameters)
            elif action in ("pdf", "to_pdf", "convert"):
                return _word_to_pdf(parameters)
            else:
                return _word_create(parameters)

        # ── EXCEL ──────────────────────────────────────────────────────────
        elif app in ("excel", "spreadsheet", "xlsx", "hoja"):
            if action in ("create", "new", "make"):
                return _excel_create(parameters)
            elif action in ("analyze", "stats", "summary", "info"):
                return _excel_analyze(parameters)
            elif action in ("add", "append", "insert"):
                return _excel_add_data(parameters)
            else:
                return _excel_create(parameters)

        # ── POWERPOINT ─────────────────────────────────────────────────────
        elif app in ("powerpoint", "pptx", "presentation", "slides"):
            if action in ("create", "new", "make", "build"):
                return _pptx_create(parameters)
            else:
                return _pptx_create(parameters)

        else:
            return f"Unknown app '{app}'. Use: word, excel, or powerpoint."

    except Exception as e:
        return f"office_controller error [{app}/{action}]: {e}"
